"""Independently validate source amounts, tool use and the real PPTX download."""
import hashlib
import json
import re
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation

root = Path(__file__).resolve().parent
directory = Path((root / 'primeagent-quality/latest.txt').read_text().strip())
load = lambda path: json.loads(path.read_text(encoding='utf-8-sig'))
plan = load(directory / 'plan.json')
browser = load(directory / 'browser-verification-after-workspace-fix.json')
checks = []

def check(name, function):
    try:
        assert function(), name
        checks.append({'name': name, 'passed': True})
    except Exception as error:
        checks.append({'name': name, 'passed': False, 'error': str(error)})

document = next(row for row in browser['attempts'] if row['task'] == 'documents')
presentation = next(row for row in browser['attempts'] if row['task'] == 'presentation')
actual = load(Path(plan['workspace']) / 'documents/invoice-result.json')
expected = load(directory / 'expected-invoice.json')
subtotal = Decimal('124.50') * 3 + Decimal('80.00') * 2
vat = (subtotal * Decimal('0.05')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
total = subtotal + vat
check('Document task completed without a turn/observer error', lambda: document['completed'] and not document.get('hasTurnError') and not document.get('observerError'))
for key in ['invoiceId', 'supplier', 'currency', 'lineTotals']:
    check('Source field ' + key, lambda key=key: actual[key] == expected[key])
for key, number in [('subtotal', subtotal), ('vat', vat), ('total', total)]:
    check('Exact amount ' + key, lambda key=key, number=number: isinstance(actual[key], (int, float)) and Decimal(str(actual[key])) == number)
check('Connected document reader actually completed', lambda: any(row['tool'] == 'mcp__local-workbench__read_document' and row['status'] == 'completed' for row in document['tools']))
check('Connected exact-math tool actually completed', lambda: any(row['tool'] == 'mcp__local-workbench__math_compute' and row['status'] == 'completed' for row in document['tools']))
check('Presentation task completed without a turn/observer error', lambda: presentation['completed'] and not presentation.get('hasTurnError') and not presentation.get('observerError'))
check('Presentation MCP tool actually completed', lambda: any(row['tool'] == 'mcp__local-workbench__create_presentation' and row['status'] == 'completed' for row in presentation['tools']))
check('Native artifact registration actually completed', lambda: any(row['tool'] == 'record_artifact' and row['status'] == 'completed' for row in presentation['tools']))
download = directory / 'downloaded-presentation.pptx'
check('Actual browser Download yielded the recorded file', lambda: bool(presentation.get('download')) and download.stat().st_size == presentation['download']['bytes'] and hashlib.sha256(download.read_bytes()).hexdigest() == presentation['download']['sha256'])
portal = load(root.parent / 'portal.config.json')
native_output = Path(portal['agentWorkspace']) / 'output' / presentation['download']['filename']
check('Downloaded bytes match the native generated artifact', lambda: native_output.read_bytes() == download.read_bytes())
with ZipFile(download) as archive:
    check('Downloaded OOXML archive is intact', lambda: archive.testzip() is None)
deck = Presentation(str(download))
check('Exactly two slides', lambda: len(deck.slides) == 2)
texts = []
for slide in deck.slides:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    texts.append('\n'.join(parts))
all_text = '\n'.join(texts)
check('PrimeAgent and Asaad AlSharif attribution', lambda: 'PrimeAgent — Asaad AlSharif' in texts[0])
check('Second slide title', lambda: 'Verified totals' in texts[1])
check('Arabic supplier preserved in presentation', lambda: expected['supplier'] in all_text)
check('Invoice ID and currency preserved', lambda: expected['invoiceId'] in all_text and expected['currency'] in all_text)
amounts = {Decimal(value) for value in re.findall(r'(?<![\d.])\d+(?:\.\d+)?(?![\d.])', texts[1])}
check('All independently calculated totals appear on slide two', lambda: {subtotal, vat, total}.issubset(amounts))
line_amounts = {Decimal(value) for value in re.findall(r'(?<![\d.])\d+(?:\.\d+)?(?![\d.])', texts[0])}
check('Both line totals appear on slide one', lambda: {Decimal('373.50'), Decimal('160.00')}.issubset(line_amounts))
result = {'checkedAtUtc': datetime.now(timezone.utc).isoformat(), 'checks': len(checks), 'passed': sum(row['passed'] for row in checks), 'failed': [row for row in checks if not row['passed']], 'details': checks, 'invoice': actual, 'download': presentation.get('download'), 'slideText': texts, 'sessionUrls': {'documents': document['sessionUrl'], 'presentation': presentation['sessionUrl']}}
(directory / 'independent-document-verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({key: result[key] for key in ['checks','passed','failed','download']}, ensure_ascii=True))
raise SystemExit(bool(result['failed']))
